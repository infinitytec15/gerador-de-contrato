from flask import Flask, render_template, request, send_file
from docx import Document
import pdfkit
from io import BytesIO
import os
import subprocess
import logging

app = Flask(__name__)

# Configurar logs
logging.basicConfig(level=logging.DEBUG)
logger = logging.getLogger(__name__)

# Configuração do pdfkit com o caminho do wkhtmltopdf
pdfkit_config = pdfkit.configuration(wkhtmltopdf='/usr/bin/wkhtmltopdf')

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/gerar_procuracao', methods=['POST'])
def gerar_procuracao():
    try:
        # Coletar dados do formulário
        dados = {
            'nome': request.form['nome'],
            'apelido': request.form['apelido'],
            'nacionalidade': request.form['nacionalidade'],
            'profissao': request.form['profissao'],
            'estado_civil': request.form['estado_civil'],
            'rg_uf': request.form['rg_uf'],
            'cpf': request.form['cpf'],
            'rua': request.form['rua'],
            'numero': request.form['numero'],
            'bairro': request.form['bairro'],
            'municipio_uf': request.form['municipio_uf'],
            'cep': request.form['cep'],
            'telefone': request.form['telefone']
        }

        logger.debug("Dados do formulário coletados com sucesso.")

        # Abrir o template DOCX
        doc = Document('templates/procuração_template.docx')
        logger.debug("Template DOCX carregado com sucesso.")

        # Substituir placeholders pelos dados
        for paragraph in doc.paragraphs:
            for key, value in dados.items():
                if f'{{{key}}}' in paragraph.text:
                    paragraph.text = paragraph.text.replace(f'{{{key}}}', value)

        # Salvar o DOCX preenchido temporariamente
        docx_temp_path = 'temp_procuracao.docx'
        doc.save(docx_temp_path)
        logger.debug(f"DOCX preenchido salvo em: {docx_temp_path}")

        # Converter o DOCX para HTML usando pandoc
        html_temp_path = 'temp_procuracao.html'
        command = f'pandoc "{docx_temp_path}" -o "{html_temp_path}" --from=docx --to=html'
        result = subprocess.run(command, shell=True, capture_output=True, text=True)

        # Verificar se o comando pandoc foi executado com sucesso
        if result.returncode != 0:
            logger.error(f"Erro ao converter DOCX para HTML: {result.stderr}")
            raise RuntimeError(f"Erro ao converter DOCX para HTML: {result.stderr}")

        logger.debug(f"HTML gerado salvo em: {html_temp_path}")

        # Ler o conteúdo HTML com codificação UTF-8
        with open(html_temp_path, 'r', encoding='utf-8') as f:
            html_content = f.read()

        # Converter o HTML para PDF usando pdfkit
        pdf_temp_path = 'temp_procuracao.pdf'
        pdfkit.from_string(html_content, pdf_temp_path, configuration=pdfkit_config, options={"encoding": "UTF-8"})

        logger.debug(f"PDF gerado salvo em: {pdf_temp_path}")

        # Ler o PDF gerado
        with open(pdf_temp_path, 'rb') as f:
            pdf_data = f.read()

        # Retornar o PDF para o usuário
        return send_file(
            BytesIO(pdf_data),
            as_attachment=True,
            download_name='procuração.pdf',
            mimetype='application/pdf'
        )

    except Exception as e:
        logger.error(f"Erro ao gerar o PDF: {str(e)}")
        return f"Erro ao gerar o PDF: {str(e)}", 500

    finally:
        # Remover arquivos temporários
        if os.path.exists(docx_temp_path):
            os.remove(docx_temp_path)
        if os.path.exists(html_temp_path):
            os.remove(html_temp_path)
        if os.path.exists(pdf_temp_path):
            os.remove(pdf_temp_path)
        logger.debug("Arquivos temporários removidos.")

if __name__ == '__main__':
    app.run(debug=True)